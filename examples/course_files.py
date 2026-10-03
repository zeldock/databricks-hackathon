"""Render the fictional courses in course_content.py into realistic files.

For each course this writes:
  syllabi/<slug>_Syllabus.pdf                  — full syllabus with a dated schedule and deadlines
  course_materials/<slug>/*.pptx               — lecture slide decks
  course_materials/<slug>/*.docx               — lecture notes
  course_materials/<slug>/*.pdf                — practice problems / study guide / assignment sheet

All dates are computed from the day the script runs, so every deadline is in the future.
"""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

import matplotlib
from docx import Document
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGB
from pptx.util import Inches, Pt as PptxPt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate,
                                Spacer, Table, TableStyle)

import course_content as content

ACCENTS = {"MA 141": "#1f6feb", "CSC 113": "#2da44e", "ENG 101": "#8250df", "PSY 200": "#d1242f"}
TYPE_LABEL = {"Deadline": "Assignment due", "Exam": "Exam", "Quiz": "Quiz", "Milestone": "Project milestone"}

# DejaVu Sans ships with matplotlib and covers math symbols (≤, →, ², √, π) that the built-in PDF fonts lack.
_FONT_DIR = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
pdfmetrics.registerFont(TTFont("DejaVu", str(_FONT_DIR / "DejaVuSans.ttf")))
pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(_FONT_DIR / "DejaVuSans-Bold.ttf")))
pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold", italic="DejaVu", boldItalic="DejaVu-Bold")


# --------------------------------------------------------------------------- #
# Dates
# --------------------------------------------------------------------------- #
def session_start(today: date) -> date:
    """The Monday strictly after today — so every date in the session is in the future."""
    return today + timedelta(days=7 - today.weekday())


def thanksgiving(year: int) -> date:
    first = date(year, 11, 1)
    first_thursday = first + timedelta(days=(3 - first.weekday()) % 7)
    return first_thursday + timedelta(weeks=3)


def dated_deliverables(course: dict, start: date) -> list[dict]:
    """Turn (week, weekday, title, type) into real dates, moving anything off the Thanksgiving break."""
    holiday = thanksgiving(start.year)
    out = []
    for week, weekday, title, kind in course["deliverables"]:
        when = start + timedelta(weeks=week - 1, days=weekday)
        if holiday - timedelta(days=1) <= when <= holiday + timedelta(days=3):
            # Move off the Thanksgiving break: exams to the Monday before, everything else to the Tuesday.
            when = holiday - timedelta(days=3 if kind == "Exam" else 2)
        out.append({"week": week, "date": when, "title": title, "type": kind})
    return sorted(out, key=lambda d: d["date"])


def fmt(day: date) -> str:
    return f"{day:%a}, {day:%b} {day.day}, {day.year}"


def fmt_short(day: date) -> str:
    return f"{day:%a} {day:%b} {day.day}"


# --------------------------------------------------------------------------- #
# PDF helpers
# --------------------------------------------------------------------------- #
def _styles(accent: str) -> dict[str, ParagraphStyle]:
    base = dict(fontName="DejaVu", fontSize=9.5, leading=13.5, textColor=colors.HexColor("#1f2328"))
    return {
        "uni": ParagraphStyle("uni", **{**base, "fontSize": 9, "textColor": colors.HexColor("#57606a")}),
        "title": ParagraphStyle("title", **{**base, "fontName": "DejaVu-Bold", "fontSize": 20, "leading": 25,
                                            "textColor": colors.HexColor(accent)}),
        "subtitle": ParagraphStyle("subtitle", **{**base, "fontSize": 11, "leading": 15}),
        "h2": ParagraphStyle("h2", **{**base, "fontName": "DejaVu-Bold", "fontSize": 12.5, "leading": 16,
                                      "spaceBefore": 12, "spaceAfter": 5, "textColor": colors.HexColor(accent)}),
        "body": ParagraphStyle("body", **{**base, "alignment": TA_LEFT, "spaceAfter": 5}),
        "small": ParagraphStyle("small", **{**base, "fontSize": 8.5, "leading": 11.5}),
        "cell": ParagraphStyle("cell", **{**base, "fontSize": 8.8, "leading": 11.5}),
        "cellbold": ParagraphStyle("cellbold", **{**base, "fontName": "DejaVu-Bold", "fontSize": 8.8,
                                                  "leading": 11.5}),
    }


def _bullets(items: list[str], style: ParagraphStyle) -> ListFlowable:
    return ListFlowable([ListItem(Paragraph(i, style), leftIndent=12) for i in items], bulletType="bullet",
                        start="•", leftIndent=12, bulletFontName="DejaVu")


def _table(rows: list[list], widths: list[float], accent: str, header: bool = True) -> Table:
    table = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
    style = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d0d7de")),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    if header:
        style += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(accent)),
                  ("TEXTCOLOR", (0, 0), (-1, 0), colors.white)]
        style += [("BACKGROUND", (0, r), (-1, r), colors.HexColor("#f6f8fa")) for r in range(2, len(rows), 2)]
    table.setStyle(TableStyle(style))
    return table


def _footer(text: str):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("DejaVu", 7.5)
        canvas.setFillColor(colors.HexColor("#6e7781"))
        canvas.drawString(0.75 * inch, 0.5 * inch, text)
        canvas.drawRightString(letter[0] - 0.75 * inch, 0.5 * inch, f"Page {doc.page}")
        canvas.restoreState()
    return draw


def _doc(path: Path) -> SimpleDocTemplate:
    return SimpleDocTemplate(str(path), pagesize=letter, leftMargin=0.75 * inch, rightMargin=0.75 * inch,
                             topMargin=0.7 * inch, bottomMargin=0.8 * inch)


# --------------------------------------------------------------------------- #
# Syllabus PDF
# --------------------------------------------------------------------------- #
def write_syllabus(course: dict, start: date, path: Path) -> list[dict]:
    accent = ACCENTS.get(course["code"], "#1f6feb")
    s = _styles(accent)
    deliverables = dated_deliverables(course, start)
    end = start + timedelta(weeks=10) - timedelta(days=3)  # Friday of week 10
    holiday = thanksgiving(start.year)
    cell = lambda text, bold=False: Paragraph(text, s["cellbold" if bold else "cell"])  # noqa: E731

    story = [
        Paragraph(content.UNIVERSITY.upper() + " · DEPARTMENT COURSE SYLLABUS", s["uni"]),
        Spacer(1, 4),
        Paragraph(f"{course['code']}: {course['title']}", s["title"]),
        Paragraph(f"{content.TERM_NAME} · {fmt_short(start)} – {fmt_short(end)}, {end.year} · "
                  f"{course['credits']} credit hours", s["subtitle"]),
        Spacer(1, 10),
    ]

    info = [
        ["Instructor", course["instructor"]],
        ["Email", course["email"]],
        ["Office", course["office"]],
        ["Office hours", course["office_hours"]],
        ["Class meetings", course["meets"]],
    ]
    if course.get("ta"):
        info.append(["Teaching assistants", course["ta"]])
    info.append(["Prerequisite", course["prereq"]])
    story.append(_table([[cell(k, True), cell(v)] for k, v in info], [1.45 * inch, 5.55 * inch], accent,
                        header=False))

    story += [Paragraph("Course description", s["h2"]), Paragraph(course["description"], s["body"])]
    story += [Paragraph("Learning outcomes", s["h2"]),
              Paragraph("By the end of this course you will be able to:", s["body"]),
              _bullets(course["outcomes"], s["body"])]
    story += [Paragraph("Required materials", s["h2"]), _bullets(course["materials"], s["body"])]

    grading_rows = [[cell("Component", True), cell("Weight", True)]]
    grading_rows += [[cell(name), cell(f"{weight}%")] for name, weight in course["grading"]]
    grading_rows.append([cell("Total", True), cell(f"{sum(w for _, w in course['grading'])}%", True)])
    story.append(KeepTogether([
        Paragraph("Grading", s["h2"]),
        _table(grading_rows, [5.0 * inch, 2.0 * inch], accent),
        Spacer(1, 4),
        Paragraph(f"<b>Grading scale:</b> {content.GRADE_SCALE}. Final grades are not curved; scores are rounded "
                  "to the nearest whole percent.", s["body"]),
    ]))
    story += [Paragraph("Major project", s["h2"]), Paragraph(course["project"], s["body"])]

    # Week-by-week schedule
    sched = [[cell("Week", True), cell("Dates", True), cell("Topics", True), cell("Due this week", True)]]
    for week, topic in enumerate(course["topics"], start=1):
        monday = start + timedelta(weeks=week - 1)
        friday = monday + timedelta(days=4)
        due = [f"{fmt_short(d['date'])}: {d['title']}" for d in deliverables if d["week"] == week]
        if monday <= holiday <= friday + timedelta(days=2):
            topic += f" <i>(No class {fmt_short(holiday - timedelta(days=1))} – {fmt_short(holiday + timedelta(days=1))}"\
                     ": Thanksgiving break)</i>"
        sched.append([cell(str(week)), cell(f"{monday:%b} {monday.day} – {friday:%b} {friday.day}"), cell(topic),
                      cell("<br/>".join(due) or "—")])
    story += [Paragraph("Course schedule", s["h2"]),
              Paragraph("The schedule may change with notice; changes will be announced in class and on the course "
                        "website.", s["small"]), Spacer(1, 4),
              _table(sched, [0.62 * inch, 1.13 * inch, 2.75 * inch, 2.5 * inch], accent)]

    # Every deadline in one list — this is what the dashboard's AI syllabus import reads.
    times = {"Deadline": "11:59 PM", "Milestone": "11:59 PM", "Quiz": "in class", "Exam": "in class"}
    if course["code"] == "PSY 200":
        times["Quiz"] = "11:59 PM (online)"
    dates = [[cell("Date", True), cell("Item", True), cell("Type", True), cell("Time", True)]]
    dates += [[cell(fmt(d["date"])), cell(d["title"]), cell(TYPE_LABEL[d["type"]]), cell(times[d["type"]])]
              for d in deliverables]
    story += [Paragraph("Important dates and deadlines", s["h2"]),
              _table(dates, [1.75 * inch, 3.05 * inch, 1.3 * inch, 0.9 * inch], accent)]

    story += [Paragraph("Course policies", s["h2"]),
              Paragraph(f"<b>Late work and make-ups.</b> {course['late_policy']}", s["body"])]
    story += [Paragraph(f"<b>{title}.</b> {text}", s["body"]) for title, text in content.COMMON_POLICIES]
    story += [Paragraph("Getting help", s["h2"]),
              Paragraph("Come to office hours early and often — you do not need a specific question. "
                        "Free campus resources:", s["body"]),
              _bullets(content.SUPPORT, s["body"])]

    footer = (f"{course['code']} Syllabus · {content.UNIVERSITY} · Fictional course created for demonstration "
              "purposes")
    _doc(path).build(story, onFirstPage=_footer(footer), onLaterPages=_footer(footer))
    return deliverables


# --------------------------------------------------------------------------- #
# Practice / study guide PDF
# --------------------------------------------------------------------------- #
def write_practice(course: dict, path: Path) -> None:
    accent = ACCENTS.get(course["code"], "#1f6feb")
    s = _styles(accent)
    p = course["practice"]
    story = [Paragraph(content.UNIVERSITY.upper(), s["uni"]), Spacer(1, 4),
             Paragraph(p["title"], s["title"]), Spacer(1, 6), Paragraph(p["intro"], s["subtitle"]), Spacer(1, 10)]
    story.append(ListFlowable([ListItem(Paragraph(item, s["body"]), leftIndent=16) for item in p["items"]],
                              bulletType="1", bulletFormat="%s.", bulletFontSize=9.5, leftIndent=16, bulletFontName="DejaVu"))
    if p.get("answers"):
        story += [Spacer(1, 18), Paragraph("Answers", s["h2"]),
                  ListFlowable([ListItem(Paragraph(a, s["body"]), leftIndent=16) for a in p["answers"]],
                               bulletType="1", bulletFormat="%s.", bulletFontSize=9.5, leftIndent=16, bulletFontName="DejaVu")]
    footer = f"{course['code']} · {course['title']} · {content.UNIVERSITY} (fictional)"
    _doc(path).build(story, onFirstPage=_footer(footer), onLaterPages=_footer(footer))


# --------------------------------------------------------------------------- #
# Slides (PPTX)
# --------------------------------------------------------------------------- #
def _hex(color: str) -> PptxRGB:
    return PptxRGB.from_string(color.lstrip("#"))


def write_slides(course: dict, deck: dict, path: Path) -> None:
    accent = ACCENTS.get(course["code"], "#1f6feb")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    def accent_bar(slide):
        bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.18))  # 1 = rectangle
        bar.fill.solid()
        bar.fill.fore_color.rgb = _hex(accent)
        bar.line.fill.background()

    def footer(slide, number):
        box = slide.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(12.1), Inches(0.4))
        para = box.text_frame.paragraphs[0]
        para.text = f"{course['code']} · {course['title']} · {content.UNIVERSITY}" + (f"    {number}" if number else "")
        para.runs[0].font.size = PptxPt(11)
        para.runs[0].font.color.rgb = _hex("#6e7781")

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    accent_bar(title_slide)
    title_slide.shapes.title.text = deck["title"]
    title_slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = _hex(accent)
    title_slide.placeholders[1].text = f"{deck['subtitle']}\n{course['instructor']}"
    footer(title_slide, None)

    for number, (heading, bullets) in enumerate(deck["slides"], start=2):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        accent_bar(slide)
        slide.shapes.title.text = heading
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = _hex(accent)
        body = slide.placeholders[1]
        body.left, body.top, body.width, body.height = Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.0)
        frame = body.text_frame
        frame.text = bullets[0]
        for bullet in bullets[1:]:
            frame.add_paragraph().text = bullet
        for para in frame.paragraphs:
            for run in para.runs:
                run.font.size = PptxPt(26)
        footer(slide, number)
    prs.save(str(path))


# --------------------------------------------------------------------------- #
# Notes (DOCX)
# --------------------------------------------------------------------------- #
def write_notes(course: dict, path: Path) -> None:
    notes = course["notes"]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    heading = doc.add_heading(notes["title"], level=0)
    heading.runs[0].font.color.rgb = RGBColor.from_string(ACCENTS.get(course["code"], "#1f6feb").lstrip("#"))
    meta = doc.add_paragraph(f"{course['instructor']} · {content.UNIVERSITY} (fictional course for demonstration)")
    meta.runs[0].italic = True
    for section, lines in notes["sections"]:
        doc.add_heading(section, level=1)
        for line in lines:
            if line.startswith("• "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line[:2].rstrip(".").isdigit() and line[1:3].startswith(". "):
                doc.add_paragraph(line[3:], style="List Number")
            elif "\n" in line:  # code sample
                para = doc.add_paragraph()
                run = para.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(10)
            else:
                doc.add_paragraph(line)
    doc.save(str(path))


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #
def write_all(out: Path, today: date) -> list[Path]:
    start = session_start(today)
    written: list[Path] = []
    syllabi = out / "syllabi"
    syllabi.mkdir(parents=True, exist_ok=True)
    for course in content.COURSES:
        assert sum(w for _, w in course["grading"]) == 100, course["code"]
        syllabus = syllabi / f"{course['slug']}_Syllabus.pdf"
        deliverables = write_syllabus(course, start, syllabus)
        assert all(d["date"] > today for d in deliverables)
        written.append(syllabus)

        folder = out / "course_materials" / course["slug"]
        folder.mkdir(parents=True, exist_ok=True)
        for deck in course["slides"]:
            write_slides(course, deck, folder / deck["file"])
            written.append(folder / deck["file"])
        write_notes(course, folder / course["notes"]["file"])
        written.append(folder / course["notes"]["file"])
        write_practice(course, folder / course["practice"]["file"])
        written.append(folder / course["practice"]["file"])
    return written
